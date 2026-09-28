"""Recalcula un .xlsx con LibreOffice y busca errores de fórmula.
Uso: python3 verificar_xlsx.py archivo.xlsx"""
import os, shutil, subprocess, sys, tempfile
import openpyxl

ruta = sys.argv[1]
tmp = tempfile.mkdtemp()
shutil.copy(ruta, os.path.join(tmp, "in.xlsx"))
subprocess.run(["soffice", "--headless", "--calc", "--convert-to", "xlsx:Calc MS Excel 2007 XML",
                "--outdir", os.path.join(tmp, "out"), os.path.join(tmp, "in.xlsx")],
               capture_output=True, timeout=180)
rec = os.path.join(tmp, "out", "in.xlsx")
if not os.path.exists(rec):
    print("   ❌ LibreOffice no pudo recalcular"); sys.exit(1)
wb_f = openpyxl.load_workbook(ruta)                  # fórmulas
wb_v = openpyxl.load_workbook(rec, data_only=True)   # valores recalculados
errores, n_formulas = [], 0
for ws in wb_f.worksheets:
    wv = wb_v[ws.title]
    for row in ws.iter_rows():
        for c in row:
            if isinstance(c.value, str) and c.value.startswith("="):
                n_formulas += 1
                v = wv[c.coordinate].value
                if v is None or (isinstance(v, str) and v.startswith("#")) or (isinstance(v, str) and "Err:" in v):
                    errores.append(f"{ws.title}!{c.coordinate} {c.value[:60]} -> {v}")
print(f"   hojas={wb_f.sheetnames} fórmulas={n_formulas}")
if errores:
    print("   ❌ errores de fórmula (primeros 20):"); print("\n".join("     " + e for e in errores[:20])); sys.exit(1)
print("   ✓ planilla recalcula sin errores")
